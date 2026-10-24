# app.py
from flask import Flask, render_template, request, jsonify, send_file
from datetime import datetime
from threading import Lock
from io import BytesIO
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from database import (
    initialize_database,
    save_sensor_reading,
    get_sensor_history,
    save_pump_log
)

app = Flask(__name__)

DEVICE_TOKEN = os.environ.get("DEVICE_TOKEN", "").strip()
data_lock = Lock()

sensor_data = {
    "soil1_raw": 0,
    "soil2_raw": 0,
    "soil1_moisture": 0,
    "soil2_moisture": 0,
    "rain_raw": 0,
    "rain_detected": False,
    "flow_pulses": 0,
    "water_liters": 0,
    "pump_status": False,
    "flow_fault": False,
    "runtime_fault": False,
    "automatic_mode": True,
    "last_update": None,
    "last_seen": None
}

control_state = {"mode": "AUTO", "pump": 0}


def authorized():
    if not DEVICE_TOKEN:
        return True

    token = request.headers.get("Authorization", "")
    return token == f"Bearer {DEVICE_TOKEN}"


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/health")
def health():
    return jsonify({
        "status": "online",
        "message": "Smart Irrigation API is running"
    })


@app.route("/api/sensors", methods=["POST"])
def receive_sensors():
    if not authorized():
        return jsonify({"error": "Unauthorized device"}), 401

    payload = request.get_json(silent=True)

    if not isinstance(payload, dict):
        return jsonify({"error": "Invalid JSON data"}), 400

    required_fields = [
        "soil1_raw", "soil2_raw",
        "soil1_moisture", "soil2_moisture",
        "rain_raw", "rain_detected",
        "flow_pulses", "water_liters",
        "pump_status", "flow_fault",
        "runtime_fault"
    ]

    for field in required_fields:
        if field not in payload:
            return jsonify({"error": f"Missing field: {field}"}), 400

    try:
        clean_data = {
            "soil1_raw": int(payload["soil1_raw"]),
            "soil2_raw": int(payload["soil2_raw"]),
            "soil1_moisture": float(payload["soil1_moisture"]),
            "soil2_moisture": float(payload["soil2_moisture"]),
            "rain_raw": int(payload["rain_raw"]),
            "rain_detected": payload["rain_detected"],
            "flow_pulses": int(payload["flow_pulses"]),
            "water_liters": float(payload["water_liters"]),
            "pump_status": payload["pump_status"],
            "flow_fault": payload["flow_fault"],
            "runtime_fault": payload["runtime_fault"]
        }

        for field in [
            "rain_detected", "pump_status",
            "flow_fault", "runtime_fault"
        ]:
            if not isinstance(clean_data[field], bool):
                raise ValueError(f"{field} must be true or false")

    except (TypeError, ValueError):
        return jsonify({"error": "Invalid sensor values"}), 400

    if (
        clean_data["soil1_raw"] < 0
        or clean_data["soil2_raw"] < 0
        or clean_data["rain_raw"] < 0
        or clean_data["flow_pulses"] < 0
        or clean_data["water_liters"] < 0
        or not 0 <= clean_data["soil1_moisture"] <= 100
        or not 0 <= clean_data["soil2_moisture"] <= 100
    ):
        return jsonify({"error": "Sensor value out of range"}), 400

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with data_lock:
        previous_pump_status = sensor_data["pump_status"]
        sensor_data.update(clean_data)
        sensor_data["automatic_mode"] = control_state["mode"] == "AUTO"
        sensor_data["last_update"] = now
        sensor_data["last_seen"] = now
        current_mode = control_state["mode"]

    saved = save_sensor_reading(clean_data)

    if previous_pump_status != clean_data["pump_status"]:
        save_pump_log(
            clean_data["pump_status"],
            current_mode,
            "Pump status changed"
        )

    return jsonify({
        "status": "success",
        "database_saved": saved,
        "message": "Sensor data received"
    })


@app.route("/api/sensors", methods=["GET"])
def get_sensors():
    with data_lock:
        result = sensor_data.copy()
        mode = control_state["mode"]

    device_online = False
    last_seen = result.get("last_seen")

    if last_seen:
        try:
            last_seen_time = datetime.strptime(
                last_seen, "%Y-%m-%d %H:%M:%S"
            )
            age_seconds = (
                datetime.now() - last_seen_time
            ).total_seconds()
            device_online = 0 <= age_seconds <= 30
        except ValueError:
            device_online = False

    result["device_online"] = device_online
    result["mode"] = mode

    return jsonify(result)


@app.route("/api/control", methods=["GET"])
def get_control():
    if not authorized():
        return jsonify({"error": "Unauthorized device"}), 401

    with data_lock:
        return jsonify({
            "mode": control_state["mode"],
            "pump": int(control_state["pump"])
        })


@app.route("/api/control", methods=["POST"])
def update_control():
    payload = request.get_json(silent=True)

    if not isinstance(payload, dict):
        return jsonify({
            "status": "error",
            "error": "Invalid JSON data"
        }), 400

    mode = payload.get("mode")
    pump = payload.get("pump")

    if mode is not None and mode not in ("AUTO", "MANUAL"):
        return jsonify({
            "status": "error",
            "error": "Mode must be AUTO or MANUAL"
        }), 400

    if pump is not None:
        try:
            if isinstance(pump, bool):
                raise ValueError

            pump_number = int(pump)

            if str(pump).strip() not in (
                str(pump_number), f"{pump_number}.0"
            ):
                raise ValueError

            pump = pump_number

        except (TypeError, ValueError):
            return jsonify({
                "status": "error",
                "error": "Pump must be 0 or 1"
            }), 400

        if pump not in (0, 1):
            return jsonify({
                "status": "error",
                "error": "Pump must be 0 or 1"
            }), 400

    with data_lock:
        old_mode = control_state["mode"]
        old_pump = control_state["pump"]

        if mode is not None:
            control_state["mode"] = mode

        if control_state["mode"] == "AUTO":
            control_state["pump"] = 0
        elif pump is not None:
            control_state["pump"] = pump

        new_mode = control_state["mode"]
        new_pump = control_state["pump"]

        sensor_data["automatic_mode"] = new_mode == "AUTO"

    if old_mode != new_mode or old_pump != new_pump:
        save_pump_log(
            bool(new_pump),
            new_mode,
            "Dashboard control command"
        )

    return jsonify({
        "status": "success",
        "control": {
            "mode": new_mode,
            "pump": new_pump
        }
    })


@app.route("/download-excel", methods=["GET"])
def download_excel():
    readings = get_sensor_history(limit=1000)

    if not readings:
        return jsonify({
            "error": "No sensor data found in database"
        }), 404

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sensor Readings"

    headers = [
        "ID", "Soil 1 Raw", "Soil 2 Raw",
        "Soil 1 Moisture (%)", "Soil 2 Moisture (%)",
        "Rain Raw", "Rain Detected", "Flow Pulses",
        "Water Liters", "Pump Status", "Flow Fault",
        "Runtime Fault", "Recorded At"
    ]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1976D2"
        )
        cell.alignment = Alignment(horizontal="center")

    for reading in reversed(readings):
        sheet.append([
            reading.get("id"),
            reading.get("soil1_raw"),
            reading.get("soil2_raw"),
            reading.get("soil1_moisture"),
            reading.get("soil2_moisture"),
            reading.get("rain_raw"),
            reading.get("rain_detected"),
            reading.get("flow_pulses"),
            reading.get("water_liters"),
            reading.get("pump_status"),
            reading.get("flow_fault"),
            reading.get("runtime_fault"),
            str(reading.get("recorded_at", ""))
        ])

    for column_cells in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        sheet.column_dimensions[
            column_cells[0].column_letter
        ].width = min(max_length + 2, 30)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        "smart_irrigation_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


if __name__ == "__main__":
    initialize_database()
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False
    )